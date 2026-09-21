from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal

from fastapi import FastAPI, File, HTTPException, UploadFile, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sqlalchemy import Boolean, DateTime, Integer, String, Text, create_engine, delete, select, UniqueConstraint
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
UPLOAD_DIR = ROOT / "uploads"
FRONTEND_DIR = ROOT / "frontend"
UPLOAD_DIR.mkdir(exist_ok=True)

BUNDLED_DB_PATH = DATA_DIR / "plm_demo.db"

# Deployment-safe database selection. DATABASE_URL remains the highest-precedence
# override (for PostgreSQL or an explicit SQLite URL). PLM_DB_PATH is intended
# for Railway persistent volumes such as /data/plm_demo.db. On first start only,
# the prepared demo database is copied into the persistent location. Existing
# persisted data is never overwritten.
_explicit_database_url = os.getenv("DATABASE_URL")
_runtime_db_path = os.getenv("PLM_DB_PATH")
if _explicit_database_url:
    DATABASE_URL = _explicit_database_url
elif _runtime_db_path:
    runtime_path = Path(_runtime_db_path).expanduser()
    runtime_path.parent.mkdir(parents=True, exist_ok=True)
    if not runtime_path.exists() and BUNDLED_DB_PATH.exists():
        shutil.copy2(BUNDLED_DB_PATH, runtime_path)
    DATABASE_URL = f"sqlite:///{runtime_path.as_posix()}"
else:
    DATABASE_URL = f"sqlite:///{BUNDLED_DB_PATH.as_posix()}"

connect_args = {"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}
engine = create_engine(DATABASE_URL, future=True, connect_args=connect_args)

class Base(DeclarativeBase):
    pass

class AppState(Base):
    __tablename__ = "app_state"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, default=1)
    state_json: Mapped[str] = mapped_column(Text, nullable=False)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

class SeedState(Base):
    __tablename__ = "seed_state"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, default=1)
    state_json: Mapped[str] = mapped_column(Text, nullable=False)

class ModuleState(Base):
    __tablename__ = "module_state"
    module_name: Mapped[str] = mapped_column(String(64), primary_key=True)
    data_json: Mapped[str] = mapped_column(Text, nullable=False, default="{}")
    revision: Mapped[int] = mapped_column(Integer, nullable=False, default=1)
    updated_by: Mapped[str | None] = mapped_column(String(255), nullable=True)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

class EntityRecord(Base):
    __tablename__ = "entity_records"
    entity_type: Mapped[str] = mapped_column(String(64), primary_key=True)
    entity_id: Mapped[str] = mapped_column(String(128), primary_key=True)
    name: Mapped[str | None] = mapped_column(String(255), nullable=True)
    status: Mapped[str | None] = mapped_column(String(128), nullable=True)
    domain: Mapped[str | None] = mapped_column(String(128), nullable=True)
    request_id: Mapped[str | None] = mapped_column(String(128), nullable=True)
    raw_json: Mapped[str] = mapped_column(Text, nullable=False)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

class AuditEvent(Base):
    __tablename__ = "audit_events"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    event_key: Mapped[str] = mapped_column(String(128), unique=True, index=True)
    event_type: Mapped[str | None] = mapped_column(String(128), nullable=True)
    entity_id: Mapped[str | None] = mapped_column(String(128), nullable=True)
    actor: Mapped[str | None] = mapped_column(String(255), nullable=True)
    environment: Mapped[str | None] = mapped_column(String(64), nullable=True)
    raw_json: Mapped[str] = mapped_column(Text, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

class ImportBatch(Base):
    __tablename__ = "import_batches"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    batch_id: Mapped[str] = mapped_column(String(64), unique=True, index=True)
    import_type: Mapped[str] = mapped_column(String(64))
    file_name: Mapped[str] = mapped_column(String(255))
    row_count: Mapped[int] = mapped_column(Integer, default=0)
    valid_count: Mapped[int] = mapped_column(Integer, default=0)
    error_count: Mapped[int] = mapped_column(Integer, default=0)
    status: Mapped[str] = mapped_column(String(64), default="Parsed")
    raw_json: Mapped[str] = mapped_column(Text, nullable=False)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

class Attachment(Base):
    __tablename__ = "attachments"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    entity_type: Mapped[str] = mapped_column(String(64))
    entity_id: Mapped[str] = mapped_column(String(128))
    file_name: Mapped[str] = mapped_column(String(255))
    stored_path: Mapped[str] = mapped_column(Text)
    mime_type: Mapped[str | None] = mapped_column(String(128), nullable=True)
    uploaded_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))

Base.metadata.create_all(engine)

SEED_FILE = DATA_DIR / "seed_state.json"

ENTITY_MAPPINGS: dict[str, Any] = {
    "products": lambda state: state.get("products", []),
    "templates": lambda state: state.get("templates", []),
    "requests": lambda state: state.get("requests", []),
    "users": lambda state: state.get("users", []),
    "exceptions": lambda state: state.get("exceptions", []),
    "systems": lambda state: state.get("integrations", []) + state.get("systems", []),
    "dependencies": lambda state: state.get("dependencies", []),
    "retirements": lambda state: [r for r in state.get("requests", []) if r.get("typeKey") == "retire"],
}

# Module-level persistence is the primary backend contract for demo v3.
# Each module owns a non-overlapping set of DATA branches from the frontend.
MODULE_BRANCHES: dict[str, list[str]] = {
    "catalogue": ["products", "dependencies"],
    "templates": ["templates"],
    "requests": ["requests", "impactReviews", "controlOverrides"],
    "workflow": ["approvalWorkflowV37"],
    "testing_release": ["releasePipeline", "testDataSets", "deployTargets"],
    "reconciliation": ["systemSync", "discrepanciesV37", "driftRuns", "externalChanges"],
    "exceptions": ["exceptions", "rollbackHistory"],
    "retirement": ["retirementPlans", "retirementCandidates", "legalHolds"],
    "administration": [
        "users", "delegations", "accessReviews", "adminChanges", "adminConfigV31",
        "guardrailRules", "integrations", "integrationChecks", "notificationRules",
        "notificationPreferences", "riskModelV37", "taxonomyAssignments", "baseline",
        "baselineHistory", "baselinePublished", "baselineVersion"
    ],
    "audit": ["audit"],
    "notifications": ["notifications"],
    "onboarding": ["onboardingBatches"],
}
BRANCH_TO_MODULE = {branch: module for module, branches in MODULE_BRANCHES.items() for branch in branches}

def module_payload_from_state(module_name: str, state: dict[str, Any]) -> dict[str, Any]:
    return {branch: state.get(branch) for branch in MODULE_BRANCHES[module_name] if branch in state}

def sync_modules_from_state(state: dict[str, Any], session: Session, actor: str = "system") -> None:
    for module_name in MODULE_BRANCHES:
        payload = module_payload_from_state(module_name, state)
        encoded = json.dumps(payload, ensure_ascii=False)
        row = session.get(ModuleState, module_name)
        if row is None:
            session.add(ModuleState(module_name=module_name, data_json=encoded, revision=1, updated_by=actor))
        elif json.loads(row.data_json or "{}") != payload:
            row.data_json = encoded
            row.revision += 1
            row.updated_by = actor
            row.updated_at = datetime.now(timezone.utc)

def compose_state_from_modules(session: Session, fallback: dict[str, Any] | None = None) -> dict[str, Any]:
    state = dict(fallback or {})
    rows = session.scalars(select(ModuleState)).all()
    for row in rows:
        try:
            state.update(json.loads(row.data_json or "{}"))
        except Exception:
            continue
    return state

def write_audit(session: Session, event_type: str, entity_id: str, actor: str, metadata: dict[str, Any] | None = None, environment: str = "PLM") -> None:
    raw = {
        "id": f"AUD-API-{datetime.now().strftime('%Y%m%d%H%M%S%f')}",
        "event": event_type, "ref": entity_id, "actor": actor, "environment": environment,
        "timestamp": now_iso(), "metadata": metadata or {},
    }
    session.add(AuditEvent(
        event_key=raw["id"], event_type=event_type, entity_id=entity_id, actor=actor,
        environment=environment, raw_json=json.dumps(raw, ensure_ascii=False),
    ))


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()

def stable_entity_id(item: dict[str, Any], idx: int) -> str:
    for key in ("id", "code", "email", "name"):
        if item.get(key):
            return str(item[key])
    return f"ROW-{idx+1:04d}"

def entity_status(item: dict[str, Any]) -> str | None:
    return item.get("status") or item.get("state") or item.get("lifecycleState")

def entity_name(item: dict[str, Any]) -> str | None:
    return item.get("name") or item.get("title") or item.get("target")

def entity_domain(item: dict[str, Any]) -> str | None:
    return item.get("domain") or item.get("division") or (item.get("current") or {}).get("division")

def normalize_state(state: dict[str, Any], session: Session) -> None:
    # Keep the relational mirror in sync with the prototype state. This makes
    # list/detail APIs and future module-by-module migration possible without
    # forcing a complete frontend rewrite in one step.
    for etype, getter in ENTITY_MAPPINGS.items():
        session.execute(delete(EntityRecord).where(EntityRecord.entity_type == etype))
        items = getter(state) or []
        seen = set()
        for idx, item in enumerate(items):
            if not isinstance(item, dict):
                continue
            eid = stable_entity_id(item, idx)
            # Protect against duplicate legacy seed IDs.
            if eid in seen:
                eid = f"{eid}-{idx+1}"
            seen.add(eid)
            session.add(EntityRecord(
                entity_type=etype,
                entity_id=eid,
                name=entity_name(item),
                status=entity_status(item),
                domain=entity_domain(item),
                request_id=item.get("requestId") or (eid if etype == "requests" else None),
                raw_json=json.dumps(item, ensure_ascii=False),
                updated_at=datetime.now(timezone.utc),
            ))

    # Release pipeline is object keyed by Request ID.
    session.execute(delete(EntityRecord).where(EntityRecord.entity_type == "release_packages"))
    for req_id, rec in (state.get("releasePipeline") or {}).items():
        if isinstance(rec, dict):
            session.add(EntityRecord(
                entity_type="release_packages", entity_id=str(req_id),
                name=rec.get("product") or rec.get("title") or req_id,
                status=(rec.get("state") or {}).get("closed") and "Closed" or (rec.get("stage") or "Open"),
                request_id=str(req_id), raw_json=json.dumps(rec, ensure_ascii=False),
                updated_at=datetime.now(timezone.utc),
            ))

    # Rebuild searchable audit mirror.
    session.execute(delete(AuditEvent))
    for idx, event in enumerate(state.get("audit", []) or []):
        if not isinstance(event, dict):
            continue
        key = str(event.get("id") or f"AUD-{idx+1:06d}")
        session.add(AuditEvent(
            event_key=key,
            event_type=event.get("event") or event.get("action") or event.get("type"),
            entity_id=event.get("ref") or event.get("entityId") or event.get("record"),
            actor=event.get("actor") or event.get("who") or event.get("user"),
            environment=event.get("environment") or event.get("env"),
            raw_json=json.dumps(event, ensure_ascii=False),
        ))



TEMPLATE_CLASS_DEFAULTS: dict[str, dict[str, str]] = {
    "TPL-018": {"domain": "Marketing", "family": "Voice", "category": "Tariff Plans", "subcategory": "Prepaid"},
    "TPL-022": {"domain": "Marketing", "family": "Data", "category": "Data Bundles", "subcategory": "Commercial Plans"},
    "TPL-024": {"domain": "Enterprise Business", "family": "Enterprise Mobile", "category": "Tariff Plans", "subcategory": "Enterprise Plans"},
    "TPL-025": {"domain": "Enterprise Business", "family": "Data & IoT", "category": "Data Bundles", "subcategory": "Business Plans"},
    "TPL-027": {"domain": "Marketing", "family": "Loans", "category": "Loan Services", "subcategory": "Loan Products"},
    "TPL-031": {"domain": "Marketing", "family": "Roaming", "category": "Roaming Bundles", "subcategory": "Roaming Bundles"},
}

def normalize_template_records(templates: list[Any] | None) -> tuple[list[Any], bool]:
    """Repair legacy demo template records so classification cascades never lock blank inherited values."""
    changed = False
    out: list[Any] = []
    for raw in templates or []:
        if not isinstance(raw, dict):
            out.append(raw)
            continue
        t = dict(raw)
        tid = str(t.get("id") or "")
        defaults = TEMPLATE_CLASS_DEFAULTS.get(tid, {})
        if t.get("domain") == "EB":
            t["domain"] = "Enterprise Business"
            changed = True
        for key in ("domain", "family", "category", "subcategory"):
            if not t.get(key) and defaults.get(key):
                t[key] = defaults[key]
                changed = True
        # Legacy scope fallback for records not in the default map.
        scope = str(t.get("scope") or "")
        if not t.get("category") and "," in scope:
            parts = [x.strip() for x in scope.split(",") if x.strip()]
            if parts:
                t["category"] = parts[0]
                changed = True
            if len(parts) > 1:
                sub = {("Data Bundles", "Commercial"): "Commercial Plans",
                       ("Data Bundles", "Business"): "Business Plans"}.get((parts[0], parts[1]), parts[1])
                t["subcategory"] = sub
                changed = True
        # Certified templates must carry the classification in their reusable snapshot.
        if str(t.get("state") or t.get("status") or "").lower() == "certified":
            cfg = dict(t.get("config") or {})
            for key in ("domain", "family", "category", "subcategory"):
                if t.get(key) and not cfg.get(key):
                    cfg[key] = t[key]
                    changed = True
            if cfg != (t.get("config") or {}):
                t["config"] = cfg
        out.append(t)
    return out, changed

def migrate_classification_state(state: dict[str, Any]) -> bool:
    templates, changed = normalize_template_records(state.get("templates"))
    if changed:
        state["templates"] = templates
    return changed

def migrate_persisted_classification() -> None:
    """One-time/idempotent startup repair for v4/v3 databases copied into this build."""
    with Session(engine) as session:
        changed_any = False
        for model in (SeedState, AppState):
            row = session.get(model, 1)
            if not row:
                continue
            try:
                data = json.loads(row.state_json or "{}")
            except Exception:
                continue
            if migrate_classification_state(data):
                row.state_json = json.dumps(data, ensure_ascii=False)
                if hasattr(row, "updated_at"):
                    row.updated_at = datetime.now(timezone.utc)
                changed_any = True
        mod = session.get(ModuleState, "templates")
        if mod:
            try:
                payload = json.loads(mod.data_json or "{}")
            except Exception:
                payload = {}
            templates, changed = normalize_template_records(payload.get("templates"))
            if changed:
                payload["templates"] = templates
                mod.data_json = json.dumps(payload, ensure_ascii=False)
                mod.revision += 1
                mod.updated_by = "startup-migration-v5"
                mod.updated_at = datetime.now(timezone.utc)
                changed_any = True
        if changed_any:
            # Refresh the relational mirror from the repaired composed state.
            app_row = session.get(AppState, 1)
            fallback = json.loads(app_row.state_json or "{}") if app_row else {}
            repaired = compose_state_from_modules(session, fallback)
            normalize_state(repaired, session)
            session.commit()



def _is_price_band_kind(value: Any) -> bool:
    return str(value or "").lower().replace(" ", "").replace("_", "").replace("-", "") == "priceband"

def normalize_price_band_state(state: dict[str, Any]) -> bool:
    """Normalize maker-checker approved price bands and repair legacy v5 casing/scope coexistence."""
    changed = False
    baseline = state.setdefault("baseline", [])
    if not isinstance(baseline, list):
        baseline = []
        state["baseline"] = baseline
        changed = True
    for band in baseline:
        if isinstance(band, dict) and _is_price_band_kind(band.get("kind") or band.get("ruleType")) and band.get("kind") != "priceBand":
            band["kind"] = "priceBand"
            changed = True

    cfg = state.get("adminConfigV31") or {}
    rules = cfg.get("productRules") or []
    if not isinstance(rules, list):
        return changed
    for rule in rules:
        if not isinstance(rule, dict):
            continue
        if rule.get("group") != "Product validation" or rule.get("status") == "Inactive" or not _is_price_band_kind(rule.get("ruleType") or rule.get("kind")):
            continue
        # Keep the human-readable ruleType for the UI, but persist a canonical
        # machine kind so downstream validation and migrations have one stable value.
        if rule.get("kind") != "priceBand":
            rule["kind"] = "priceBand"
            changed = True
        try:
            min_v = float(rule.get("min"))
            max_v = float(rule.get("max"))
        except (TypeError, ValueError):
            continue
        rid = str(rule.get("id") or "")
        if not rid:
            continue
        base = next((b for b in baseline if isinstance(b, dict) and str(b.get("id") or "") == rid), None)
        if base is None:
            base = {"id": rid}
            baseline.append(base)
            changed = True
        for old in baseline:
            if old is base or not isinstance(old, dict):
                continue
            if _is_price_band_kind(old.get("kind") or old.get("ruleType")) and old.get("applies") == rule.get("applies"):
                try:
                    old_min = float(old.get("min")); old_max = float(old.get("max"))
                    old_numeric = True
                except (TypeError, ValueError):
                    old_numeric = False
                if not old_numeric and old.get("status") != "Superseded":
                    old["status"] = "Superseded"
                    changed = True
        desired = {
            "name": rule.get("name"), "value": rule.get("value"), "applies": rule.get("applies"),
            "kind": "priceBand", "min": min_v, "max": max_v, "unit": rule.get("unit") or "NGN", "status": "Active"
        }
        for key, value in desired.items():
            if base.get(key) != value:
                base[key] = value
                changed = True
    return changed

def _numeric_price_band(rule: dict[str, Any] | None) -> bool:
    if not isinstance(rule, dict) or not _is_price_band_kind(rule.get("ruleType") or rule.get("kind")):
        return False
    try:
        float(rule.get("min")); float(rule.get("max"))
        return rule.get("min") not in (None, "") and rule.get("max") not in (None, "")
    except (TypeError, ValueError):
        return False

def _approved_price_band_changes(state: dict[str, Any]) -> list[dict[str, Any]]:
    """Return approved/implemented numeric price-band proposals, newest first."""
    cfg = state.get("adminConfigV31") or {}
    changes = cfg.get("governanceChanges") or []
    recovered: list[dict[str, Any]] = []
    for change in changes:
        if not isinstance(change, dict) or change.get("area") != "Product Rule":
            continue
        if change.get("state") not in {"Approved", "Implemented"}:
            continue
        payload = change.get("payload") or {}
        data = payload.get("data") or {}
        if not isinstance(data, dict) or data.get("group", "Product validation") != "Product validation":
            continue
        candidate = dict(data)
        rid = str(payload.get("id") or candidate.get("id") or change.get("targetId") or "")
        if not rid:
            continue
        candidate["id"] = rid
        candidate.setdefault("group", "Product validation")
        candidate.setdefault("status", "Active")
        if not _numeric_price_band(candidate):
            continue
        candidate["kind"] = "priceBand"
        candidate["_approved_sort"] = str(change.get("approvedAt") or change.get("decisionAt") or change.get("requested") or change.get("id") or "")
        recovered.append(candidate)
    recovered.sort(key=lambda r: r.get("_approved_sort", ""), reverse=True)
    for candidate in recovered:
        candidate.pop("_approved_sort", None)
    return recovered

def _recover_price_bands_from_app_state(app_state: dict[str, Any], admin_payload: dict[str, Any]) -> bool:
    """Recover approved price-band data if an older startup reset Administration to seed.

    A valid persisted Administration rule always wins. AppState and approved
    governance history are only used to fill missing/non-numeric legacy rows.
    """
    changed = False
    source_cfg = app_state.get("adminConfigV31") or {}
    target_cfg = admin_payload.setdefault("adminConfigV31", {})
    source_rules = source_cfg.get("productRules") or []
    target_rules = target_cfg.setdefault("productRules", [])

    candidates: list[dict[str, Any]] = []
    for src in source_rules:
        if isinstance(src, dict) and _numeric_price_band(src) and src.get("status") != "Inactive":
            item = dict(src)
            item["kind"] = "priceBand"
            candidates.append(item)
    candidates.extend(_approved_price_band_changes(app_state))

    # Keep first occurrence for each id. Active AppState rules are preferred,
    # then approved history can recover ids that are otherwise missing/reset.
    seen_ids: set[str] = set()
    for src in candidates:
        rid = str(src.get("id") or "")
        if not rid or rid in seen_ids:
            continue
        seen_ids.add(rid)
        tgt = next((r for r in target_rules if isinstance(r, dict) and str(r.get("id") or "") == rid), None)
        if tgt is None:
            target_rules.append(dict(src))
            changed = True
        elif not _numeric_price_band(tgt):
            tgt.clear()
            tgt.update(dict(src))
            changed = True
        elif _is_price_band_kind(tgt.get("ruleType") or tgt.get("kind")) and tgt.get("kind") != "priceBand":
            tgt["kind"] = "priceBand"
            changed = True

    source_baseline = app_state.get("baseline") or []
    target_baseline = admin_payload.setdefault("baseline", [])
    baseline_candidates = [dict(src) for src in source_baseline if isinstance(src, dict) and _numeric_price_band(src)]

    # If an approved proposal survived but the corresponding baseline did not,
    # synthesize the canonical active baseline row from that approved rule.
    baseline_candidates.extend({
        "id": src.get("id"),
        "name": src.get("name"),
        "value": src.get("value"),
        "applies": src.get("applies"),
        "kind": "priceBand",
        "min": src.get("min"),
        "max": src.get("max"),
        "unit": src.get("unit") or "NGN",
        "status": "Active",
    } for src in _approved_price_band_changes(app_state))

    seen_baseline_ids: set[str] = set()
    for src in baseline_candidates:
        rid = str(src.get("id") or "")
        if not rid or rid in seen_baseline_ids:
            continue
        seen_baseline_ids.add(rid)
        src["kind"] = "priceBand"
        tgt = next((b for b in target_baseline if isinstance(b, dict) and str(b.get("id") or "") == rid), None)
        if tgt is None:
            target_baseline.append(dict(src))
            changed = True
        elif not _numeric_price_band(tgt):
            tgt.clear()
            tgt.update(dict(src))
            changed = True
        elif _is_price_band_kind(tgt.get("ruleType") or tgt.get("kind")) and tgt.get("kind") != "priceBand":
            tgt["kind"] = "priceBand"
            changed = True
    return changed

def migrate_persisted_price_bands() -> None:
    """Idempotent startup repair for approved price-band changes created in v5."""
    with Session(engine) as session:
        changed_any = False
        app_data: dict[str, Any] = {}
        for model in (SeedState, AppState):
            row = session.get(model, 1)
            if not row:
                continue
            try:
                data = json.loads(row.state_json or "{}")
            except Exception:
                continue
            if model is AppState:
                app_data = data
            if normalize_price_band_state(data):
                row.state_json = json.dumps(data, ensure_ascii=False)
                if hasattr(row, "updated_at"):
                    row.updated_at = datetime.now(timezone.utc)
                changed_any = True
                if model is AppState:
                    app_data = data

        mod = session.get(ModuleState, "administration")
        if mod:
            try:
                payload = json.loads(mod.data_json or "{}")
            except Exception:
                payload = {}
            recovered = _recover_price_bands_from_app_state(app_data, payload) if app_data else False
            normalized = normalize_price_band_state(payload)
            if recovered or normalized:
                mod.data_json = json.dumps(payload, ensure_ascii=False)
                mod.revision += 1
                mod.updated_by = "startup-migration-v6-price-band"
                mod.updated_at = datetime.now(timezone.utc)
                changed_any = True

        if changed_any:
            app_row = session.get(AppState, 1)
            fallback = json.loads(app_row.state_json or "{}") if app_row else {}
            repaired = compose_state_from_modules(session, fallback)
            if app_row:
                app_row.state_json = json.dumps(repaired, ensure_ascii=False)
                app_row.updated_at = datetime.now(timezone.utc)
            normalize_state(repaired, session)
            session.commit()


TEST_ROLE_MODEL_VERSION = "7.0"
TEST_ROLE_MODULES: dict[str, list[str]] = {
    "Product Manager": ["dashboard","catalogue","templates","requests","workflow","testing","recon","exceptions","retirement","reports","audit"],
    "Product Owner": ["dashboard","catalogue","templates","requests","workflow","testing","recon","exceptions","retirement","reports","audit"],
    "Pricing / MKT BI": ["dashboard","catalogue","requests","workflow","reports","audit"],
    "Finance / Tax": ["dashboard","catalogue","requests","workflow","reports","audit"],
    "Business Assurance": ["dashboard","catalogue","requests","workflow","testing","recon","exceptions","retirement","reports","audit"],
    "Customer Experience": ["dashboard","catalogue","requests","workflow","reports","audit"],
    "Technology Approver": ["dashboard","catalogue","requests","workflow","testing","recon","reports","audit"],
    "Technical Owner": ["dashboard","catalogue","requests","workflow","testing","recon","exceptions","retirement","reports","audit"],
    "Regulatory / Legal": ["dashboard","catalogue","requests","workflow","exceptions","retirement","reports","audit"],
    "Executive Sponsor": ["dashboard","requests","workflow","reports","audit"],
    "CAB / Change Manager": ["dashboard","requests","workflow","testing","reports","audit"],
    "Test Analyst": ["dashboard","catalogue","requests","testing","reports","audit"],
    "Release Manager": ["dashboard","catalogue","requests","workflow","testing","recon","exceptions","reports","audit"],
    "Reconciliation & Monitoring Analyst": ["dashboard","catalogue","requests","testing","recon","exceptions","reports","audit"],
    "Exception & Rollback Manager": ["dashboard","catalogue","requests","testing","recon","exceptions","reports","audit"],
    "Product Governance": ["dashboard","catalogue","requests","workflow","exceptions","retirement","reports","audit"],
    "Records Manager": ["dashboard","catalogue","retirement","reports","audit"],
    "Compliance Officer": ["dashboard","catalogue","requests","workflow","exceptions","retirement","reports","audit"],
    "Audit & Compliance Reviewer": ["dashboard","catalogue","requests","workflow","testing","recon","exceptions","retirement","reports","audit"],
    "Product Classification Administrator": ["dashboard","catalogue","reports","audit","admin"],
    "Platform Administrator": ["dashboard","catalogue","templates","requests","workflow","testing","recon","exceptions","retirement","reports","audit","admin"],
}

TEST_ROLE_BLUEPRINTS: list[dict[str, Any]] = [
    {"name":"Product Manager","actions":["View","Create","Clone","Modify","Bulk Change","Suspend","Retire","Propose Template","Submit Request"],"fields":{"Pricing":"Propose","Taxation":"View","Regulatory":"View","Eligibility":"Edit","Fulfilment":"View","Charging":"View"}},
    {"name":"Product Owner","actions":["View","Create","Clone","Modify","Bulk Change","Suspend","Retire","Propose Template","Submit Request"],"fields":{"Pricing":"Propose","Taxation":"View","Regulatory":"View","Eligibility":"Edit","Fulfilment":"View","Charging":"View"}},
    {"name":"Pricing / MKT BI","actions":["View","Review Requests","Approve Pricing"],"fields":{"Pricing":"Approve/Edit","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"None","Charging":"View"}},
    {"name":"Finance / Tax","actions":["View","Review Requests","Approve Finance/Tax"],"fields":{"Pricing":"Approve","Taxation":"Approve/Edit","Regulatory":"View","Eligibility":"View","Fulfilment":"None","Charging":"View"}},
    {"name":"Business Assurance","actions":["View","Approve","Approve UAT Scope","Accept UAT Evidence","Approve Exceptions","Review Impact Analysis"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Customer Experience","actions":["View","Review Requests","Approve Customer Experience"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Technology Approver","actions":["View","Review Requests","Approve Technology","Review Integrations"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"Approve/Edit","Charging":"Approve/Edit"}},
    {"name":"Technical Owner","actions":["View","Review Impact Analysis","Review Integrations","Own Technical Dependencies","Validate Technical Readiness"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"Approve/Edit","Charging":"Approve/Edit"}},
    {"name":"Regulatory / Legal","actions":["View","Review Requests","Approve Regulatory/Legal"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"Approve/Edit","Eligibility":"View","Fulfilment":"None","Charging":"None"}},
    {"name":"Executive Sponsor","actions":["View","Executive Approval"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"CAB / Change Manager","actions":["View","Approve Change/Release"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Test Analyst","actions":["View","Manage Test Plan","Generate Test Pack","Import Test Cases","Execute Tests","Raise Defects","Retest"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Release Manager","actions":["View","Configure Release","Deploy UAT","Schedule Release","Promote Production","Deploy Production"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"Edit","Charging":"Edit"}},
    {"name":"Reconciliation & Monitoring Analyst","actions":["View","Run Reconciliation","Run Drift Scan","Resolve Discrepancies","Monitor"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Exception & Rollback Manager","actions":["View","Manage Exceptions","Approve Exceptions","Execute Rollback","Resolve Rollback"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Product Governance","actions":["View","Manage Retirement","Approve Retirement","Archive Products","Manage Lifecycle"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Records Manager","actions":["View","Manage Retention","Archive Products","Manage Records"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Compliance Officer","actions":["View","Manage Legal Holds","Review Regulatory Controls","Review Control Overrides"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"Approve/Edit","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Audit & Compliance Reviewer","actions":["View","Audit","Export Evidence","Run Reports"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Product Classification Administrator","actions":["View","Configure","Configure Product Classification","Manage Ownership"],"fields":{"Pricing":"View","Taxation":"View","Regulatory":"View","Eligibility":"View","Fulfilment":"View","Charging":"View"}},
    {"name":"Platform Administrator","actions":["View","Configure","Manage Users","Manage Roles","Manage Rules","Manage Integrations"],"fields":{"Pricing":"Approve/Edit","Taxation":"Approve/Edit","Regulatory":"Approve/Edit","Eligibility":"Edit","Fulfilment":"Approve/Edit","Charging":"Approve/Edit"}},
]

def normalize_test_role_model(state: dict[str, Any]) -> bool:
    """Idempotently add the v7 scoped test-role catalogue and demo Product Manager.

    Role-combination SoD is intentionally disabled for access assignment only.
    Governed transaction/maker-checker decisions are not changed here.
    """
    changed = False
    users = state.setdefault("users", [])
    if not isinstance(users, list):
        users = []
        state["users"] = users
        changed = True

    cfg = state.get("adminConfigV31")
    first_install = not isinstance(cfg, dict) or cfg.get("testModeRoleModelVersion") != TEST_ROLE_MODEL_VERSION

    # Chiamaka is a required seeded demo persona.
    chiamaka = next((u for u in users if isinstance(u, dict) and u.get("name") == "Chiamaka Eleodi"), None)
    if chiamaka is None:
        users.append({
            "id":"USR-DEMO-CHIAMAKA","name":"Chiamaka Eleodi","email":"chiamaka.eleodi@mtn.com",
            "employeeId":"DEMO-CHIAMAKA","title":"Product Manager","department":"Product & Commercial",
            "manager":"","role":"Product Manager","roles":["Product Manager"],"domains":["All"],
            "access":"Standard","accessStatus":"Active","state":"Active","recert":"2027-03-31",
            "recertDate":"2027-03-31","fieldOverrides":{"Pricing":"Inherit","Taxation":"Inherit",
            "Regulatory":"Inherit","Eligibility":"Inherit","Fulfilment":"Inherit","Charging":"Inherit"}
        })
        changed = True

    # Mutolib is the all-role test persona. Keep the assignment in sync with
    # the complete test catalogue so newly scoped roles appear immediately.
    admin = next((u for u in users if isinstance(u, dict) and u.get("name") == "Mutolib Ottan"), None)
    if admin is not None:
        roles = list(admin.get("roles") or ([admin.get("role")] if admin.get("role") else []))
        for role_name in [bp["name"] for bp in TEST_ROLE_BLUEPRINTS]:
            if role_name not in roles:
                roles.append(role_name)
                changed = True
        admin["roles"] = roles
        admin["role"] = roles[0] if roles else "Platform Administrator"
        admin["domains"] = ["All"]
        admin["accessStatus"] = "Active"
        admin["state"] = "Active"

    # Bare seed_state.json intentionally has no full Administration configuration;
    # do not create a partial one here. The frontend constructs it on first bootstrap.
    if isinstance(cfg, dict):
        roles = cfg.setdefault("roles", [])
        if not isinstance(roles, list):
            roles = []
            cfg["roles"] = roles
            changed = True
        for idx, bp in enumerate(TEST_ROLE_BLUEPRINTS, start=1):
            role = next((r for r in roles if isinstance(r, dict) and str(r.get("name") or "").lower() == bp["name"].lower()), None)
            if role is None:
                roles.append({"id":f"ROLE-TST-{idx:03d}","name":bp["name"],"domains":["All"],"modules":list(TEST_ROLE_MODULES.get(bp["name"], ["dashboard","reports","audit"])),"actions":list(bp["actions"]),"fields":dict(bp["fields"]),"status":"Active"})
                changed = True
                continue
            if not isinstance(role.get("domains"), list) or not role.get("domains"):
                role["domains"] = ["All"]
                changed = True
            expected_modules = list(TEST_ROLE_MODULES.get(bp["name"], ["dashboard","reports","audit"]))
            if role.get("modules") != expected_modules:
                role["modules"] = expected_modules
                changed = True
            role_actions = role.setdefault("actions", [])
            for action in bp["actions"]:
                if action not in role_actions:
                    role_actions.append(action)
                    changed = True
            fields = role.setdefault("fields", {})
            for key, value in bp["fields"].items():
                if not fields.get(key):
                    fields[key] = value
                    changed = True
            if not role.get("status"):
                role["status"] = "Active"
                changed = True

        assigned_role_names = {role_name for u in users if isinstance(u, dict) for role_name in (u.get("roles") or ([u.get("role")] if u.get("role") else []))}
        filtered_roles = [r for r in roles if not (isinstance(r, dict) and r.get("name") in {"Technology", "Regulatory"} and r.get("name") not in assigned_role_names)]
        if len(filtered_roles) != len(roles):
            cfg["roles"] = roles = filtered_roles
            changed = True

        if cfg.get("testModeRoleModelVersion") != TEST_ROLE_MODEL_VERSION:
            cfg["testModeRoleModelVersion"] = TEST_ROLE_MODEL_VERSION
            cfg["enforceRoleCombinationSoD"] = False
            cfg["accessAssignmentMode"] = "Direct test-mode assignment"
            changed = True
        if cfg.get("roleScopeVersion") != TEST_ROLE_MODEL_VERSION:
            cfg["roleScopeVersion"] = TEST_ROLE_MODEL_VERSION
            changed = True
        scopes = cfg.setdefault("userScopes", {})
        for u in users:
            if not isinstance(u, dict) or not u.get("name"):
                continue
            uroles = list(u.get("roles") or ([u.get("role")] if u.get("role") else []))
            if uroles != u.get("roles"):
                u["roles"] = uroles
                changed = True
            if not u.get("domains"):
                u["domains"] = ["All"]
                changed = True
            scopes[u["name"]] = {
                "role": uroles[0] if uroles else u.get("role", ""),
                "roles": list(uroles), "domains": list(u.get("domains") or ["All"]),
                "status": u.get("accessStatus") or u.get("state") or "Active",
                "fieldOverrides": dict(u.get("fieldOverrides") or {}),
            }

        # Resolve legacy/composite labels to concrete roles in the test model.
        aliases = {
            "Regulatory":"Regulatory / Legal", "Legal / Regulatory":"Regulatory / Legal",
            "CAB":"CAB / Change Manager", "Records / Governance":"Records Manager",
            "Compliance":"Compliance Officer", "Release / Test Analyst":"Test Analyst",
            "Approver":"Business Assurance",
        }
        for route in cfg.get("approvalRoutes") or []:
            for key in ("business", "production"):
                vals = route.get(key) or []
                new_vals = [aliases.get(v, v) for v in vals]
                if new_vals != vals:
                    route[key] = new_vals; changed = True
            if route.get("uat") in aliases:
                route["uat"] = aliases[route["uat"]]; changed = True
        for transition in cfg.get("transitions") or []:
            if transition.get("role") in aliases:
                transition["role"] = aliases[transition["role"]]; changed = True

    # Persisted approval-workflow stage labels are normalized too.
    aliases = {
        "Regulatory":"Regulatory / Legal", "Legal / Regulatory":"Regulatory / Legal",
        "CAB":"CAB / Change Manager", "Records / Governance":"Records Manager",
        "Compliance":"Compliance Officer", "Release / Test Analyst":"Test Analyst",
        "Approver":"Business Assurance",
    }
    for workflow in state.get("approvalWorkflowV37") or []:
        for stage in workflow.get("stages") or []:
            vals = stage.get("roles") or []
            new_vals = [aliases.get(v, v) for v in vals]
            if new_vals != vals:
                stage["roles"] = new_vals; changed = True
    return changed

def migrate_persisted_test_role_model() -> None:
    with Session(engine) as session:
        changed_any = False
        for model in (SeedState, AppState):
            row = session.get(model, 1)
            if not row:
                continue
            try:
                data = json.loads(row.state_json or "{}")
            except Exception:
                continue
            if normalize_test_role_model(data):
                row.state_json = json.dumps(data, ensure_ascii=False)
                if hasattr(row, "updated_at"):
                    row.updated_at = datetime.now(timezone.utc)
                changed_any = True
        mod = session.get(ModuleState, "administration")
        if mod:
            try:
                payload = json.loads(mod.data_json or "{}")
            except Exception:
                payload = {}
            if normalize_test_role_model(payload):
                mod.data_json = json.dumps(payload, ensure_ascii=False)
                mod.revision += 1
                mod.updated_by = "startup-migration-v7-scoped-roles"
                mod.updated_at = datetime.now(timezone.utc)
                changed_any = True
        if changed_any:
            app_row = session.get(AppState, 1)
            fallback = json.loads(app_row.state_json or "{}") if app_row else {}
            repaired = compose_state_from_modules(session, fallback)
            if app_row:
                app_row.state_json = json.dumps(repaired, ensure_ascii=False)
                app_row.updated_at = datetime.now(timezone.utc)
            normalize_state(repaired, session)
            session.commit()

def load_seed() -> dict[str, Any]:
    if not SEED_FILE.exists():
        return {}
    return json.loads(SEED_FILE.read_text(encoding="utf-8"))

def ensure_seed_rows() -> None:
    seed = load_seed()
    with Session(engine) as session:
        if not session.get(SeedState, 1):
            session.add(SeedState(id=1, state_json=json.dumps(seed, ensure_ascii=False)))
        if not session.get(AppState, 1):
            session.add(AppState(id=1, state_json=json.dumps(seed, ensure_ascii=False)))
            normalize_state(seed, session)
        # IMPORTANT: seed data initializes missing modules only. Existing module rows
        # are persisted user/demo state and must never be overwritten at startup.
        for module_name in MODULE_BRANCHES:
            if session.get(ModuleState, module_name) is None:
                payload = module_payload_from_state(module_name, seed)
                session.add(ModuleState(
                    module_name=module_name,
                    data_json=json.dumps(payload, ensure_ascii=False),
                    revision=1,
                    updated_by="seed",
                ))
        session.commit()

ensure_seed_rows()
migrate_persisted_classification()
migrate_persisted_price_bands()
migrate_persisted_test_role_model()

class StatePayload(BaseModel):
    data: dict[str, Any]

class ModulePayload(BaseModel):
    data: dict[str, Any]
    expected_revision: int | None = None
    actor: str | None = None

class RecordPayload(BaseModel):
    data: dict[str, Any]
    actor: str | None = None

class ApprovalDecisionPayload(BaseModel):
    role: str
    actor: str
    decision: Literal["Approved", "Rejected", "Returned"]
    note: str | None = None

class TestExecutionPayload(BaseModel):
    actor: str
    actual_result: str
    result: Literal["Pass", "Fail", "Blocked"]
    evidence: list[str] = []
    reference: str | None = None
    comments: str | None = None

class LoginPayload(BaseModel):
    email: str

class MockDeployPayload(BaseModel):
    request_id: str
    package_id: str | None = None
    result: str = "success"
    message: str | None = None

app = FastAPI(title="PLM Demo API", version="7.2.1")

@app.middleware("http")
async def disable_frontend_cache(request: Request, call_next):
    response = await call_next(request)
    if request.url.path in ("/", "/index.html"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response

@app.get("/api/health")
def health():
    return {"status": "ok", "version": "7.2.1", "database": DATABASE_URL.split(":", 1)[0], "time": now_iso()}

@app.get("/api/meta")
def meta():
    return {
        "name": "PLM Requirements Validation Demo",
        "version": "7.2.1",
        "persistence": "module-level database-backed",
        "capabilities": [
            "module-level persistence", "revision-aware APIs", "transactional CRUD", "workflow decisions",
            "test execution APIs", "roles/users", "CSV/XLSX import", "mock authoritative-system deployment",
            "audit mirror", "classification migration", "demo reset"
        ],
    }

@app.get("/api/state")
def get_state():
    with Session(engine) as session:
        row = session.get(AppState, 1)
        return {"data": json.loads(row.state_json) if row else {}}

@app.post("/api/state/bootstrap")
def bootstrap_state(payload: StatePayload):
    with Session(engine) as session:
        row = session.get(AppState, 1)
        # First browser bootstrap enriches the base seed with all derived data
        # created by the prototype's setup functions. Do not overwrite an
        # already-enriched state unless it is still the untouched base seed.
        current = json.loads(row.state_json) if row else {}
        should_enrich = len(payload.data.keys()) > len(current.keys()) or not current
        if should_enrich:
            data = payload.data
            if row:
                row.state_json = json.dumps(data, ensure_ascii=False)
                row.updated_at = datetime.now(timezone.utc)
            else:
                session.add(AppState(id=1, state_json=json.dumps(data, ensure_ascii=False)))
            # Store the enriched demo seed as the reset baseline the first time.
            seed_row = session.get(SeedState, 1)
            if seed_row and len(data.keys()) > len(json.loads(seed_row.state_json).keys()):
                seed_row.state_json = json.dumps(data, ensure_ascii=False)
            normalize_state(data, session)
            sync_modules_from_state(data, session, actor="frontend-bootstrap")
            session.commit()
            return {"data": data, "source": "frontend-enriched"}
        return {"data": current, "source": "database"}

@app.put("/api/state")
def put_state(payload: StatePayload):
    with Session(engine) as session:
        row = session.get(AppState, 1)
        encoded = json.dumps(payload.data, ensure_ascii=False)
        if row:
            row.state_json = encoded
            row.updated_at = datetime.now(timezone.utc)
        else:
            session.add(AppState(id=1, state_json=encoded))
        normalize_state(payload.data, session)
        sync_modules_from_state(payload.data, session, actor="legacy-state-sync")
        session.commit()
    return {"ok": True, "updated_at": now_iso(), "compatibility_mode": True}

@app.post("/api/demo/reset")
def demo_reset():
    with Session(engine) as session:
        seed_row = session.get(SeedState, 1)
        if not seed_row:
            raise HTTPException(500, "Seed state is unavailable")
        data = json.loads(seed_row.state_json)
        state_row = session.get(AppState, 1)
        if state_row:
            state_row.state_json = seed_row.state_json
            state_row.updated_at = datetime.now(timezone.utc)
        else:
            session.add(AppState(id=1, state_json=seed_row.state_json))
        normalize_state(data, session)
        # Reset module rows to the prepared seed as well.
        session.execute(delete(ModuleState))
        sync_modules_from_state(data, session, actor="demo-reset")
        session.commit()
    return {"ok": True, "data": data}


@app.get("/api/modules")
def list_modules():
    with Session(engine) as session:
        rows = session.scalars(select(ModuleState).order_by(ModuleState.module_name)).all()
        return [{"module": r.module_name, "revision": r.revision, "updated_by": r.updated_by,
                 "updated_at": r.updated_at.isoformat() if r.updated_at else None,
                 "branches": MODULE_BRANCHES.get(r.module_name, [])} for r in rows]

@app.get("/api/modules/bootstrap")
def bootstrap_modules():
    with Session(engine) as session:
        rows = session.scalars(select(ModuleState)).all()
        return {"modules": {r.module_name: {"revision": r.revision, "data": json.loads(r.data_json or "{}")}
                            for r in rows}, "branch_map": BRANCH_TO_MODULE}

@app.get("/api/modules/{module_name}")
def get_module(module_name: str):
    if module_name not in MODULE_BRANCHES:
        raise HTTPException(404, "Unknown module")
    with Session(engine) as session:
        row = session.get(ModuleState, module_name)
        if not row:
            return {"module": module_name, "revision": 0, "data": {}}
        return {"module": module_name, "revision": row.revision, "data": json.loads(row.data_json or "{}")}

@app.put("/api/modules/{module_name}")
def put_module(module_name: str, payload: ModulePayload):
    if module_name not in MODULE_BRANCHES:
        raise HTTPException(404, "Unknown module")
    allowed = set(MODULE_BRANCHES[module_name])
    unexpected = set(payload.data) - allowed
    if unexpected:
        raise HTTPException(400, f"Branches not owned by {module_name}: {sorted(unexpected)}")
    actor = payload.actor or "prototype-user"
    with Session(engine) as session:
        row = session.get(ModuleState, module_name)
        current_revision = row.revision if row else 0
        if payload.expected_revision is not None and payload.expected_revision != current_revision:
            raise HTTPException(409, detail={"message": "Module revision conflict", "current_revision": current_revision})
        encoded = json.dumps(payload.data, ensure_ascii=False)
        if row:
            row.data_json = encoded
            row.revision += 1
            row.updated_by = actor
            row.updated_at = datetime.now(timezone.utc)
        else:
            row = ModuleState(module_name=module_name, data_json=encoded, revision=1, updated_by=actor)
            session.add(row)
        # Compose compatibility state for old screens/endpoints, then refresh mirrors.
        app_row = session.get(AppState, 1)
        fallback = json.loads(app_row.state_json) if app_row else {}
        state = compose_state_from_modules(session, fallback)
        if app_row:
            app_row.state_json = json.dumps(state, ensure_ascii=False)
            app_row.updated_at = datetime.now(timezone.utc)
        else:
            session.add(AppState(id=1, state_json=json.dumps(state, ensure_ascii=False)))
        normalize_state(state, session)
        write_audit(session, "Module configuration persisted", module_name, actor,
                    {"revision": row.revision, "branches": list(payload.data)})
        session.commit()
        return {"ok": True, "module": module_name, "revision": row.revision, "updated_at": now_iso()}

RESOURCE_BRANCHES = {
    "products": ("catalogue", "products"), "dependencies": ("catalogue", "dependencies"),
    "templates": ("templates", "templates"), "requests": ("requests", "requests"),
    "impact-reviews": ("requests", "impactReviews"), "overrides": ("requests", "controlOverrides"),
    "users": ("administration", "users"), "delegations": ("administration", "delegations"),
    "admin-changes": ("administration", "adminChanges"), "integrations": ("administration", "integrations"),
    "notification-rules": ("administration", "notificationRules"), "notifications": ("notifications", "notifications"),
    "exceptions": ("exceptions", "exceptions"), "retirement-plans": ("retirement", "retirementPlans"),
    "onboarding-batches": ("onboarding", "onboardingBatches"), "test-data-sets": ("testing_release", "testDataSets"),
}

def record_id(item: dict[str, Any], idx: int = 0) -> str:
    return str(item.get("id") or item.get("code") or item.get("email") or item.get("name") or f"ROW-{idx+1}")

def load_module_data(session: Session, module_name: str) -> tuple[ModuleState, dict[str, Any]]:
    row = session.get(ModuleState, module_name)
    if not row:
        row = ModuleState(module_name=module_name, data_json="{}", revision=1, updated_by="system")
        session.add(row)
        session.flush()
    return row, json.loads(row.data_json or "{}")

def save_module_data(session: Session, row: ModuleState, data: dict[str, Any], actor: str, event: str, entity_id: str) -> None:
    row.data_json = json.dumps(data, ensure_ascii=False)
    row.revision += 1
    row.updated_by = actor
    row.updated_at = datetime.now(timezone.utc)
    app_row = session.get(AppState, 1)
    fallback = json.loads(app_row.state_json) if app_row else {}
    state = compose_state_from_modules(session, fallback)
    if app_row:
        app_row.state_json = json.dumps(state, ensure_ascii=False)
        app_row.updated_at = datetime.now(timezone.utc)
    normalize_state(state, session)
    write_audit(session, event, entity_id, actor)

@app.get("/api/v2/{resource}")
def v2_list(resource: str):
    if resource not in RESOURCE_BRANCHES:
        raise HTTPException(404, "Unknown resource")
    module_name, branch = RESOURCE_BRANCHES[resource]
    with Session(engine) as session:
        row, data = load_module_data(session, module_name)
        return {"items": data.get(branch, []) or [], "revision": row.revision}

@app.post("/api/v2/{resource}")
def v2_create(resource: str, payload: RecordPayload):
    if resource not in RESOURCE_BRANCHES:
        raise HTTPException(404, "Unknown resource")
    module_name, branch = RESOURCE_BRANCHES[resource]
    actor = payload.actor or "prototype-user"
    with Session(engine) as session:
        row, data = load_module_data(session, module_name)
        items = list(data.get(branch, []) or [])
        rid = record_id(payload.data, len(items))
        if any(record_id(x, i) == rid for i, x in enumerate(items) if isinstance(x, dict)):
            raise HTTPException(409, "Record already exists")
        items.append(payload.data)
        data[branch] = items
        save_module_data(session, row, data, actor, f"{resource} created", rid)
        session.commit()
        return {"item": payload.data, "revision": row.revision}

@app.get("/api/v2/{resource}/{item_id}")
def v2_get(resource: str, item_id: str):
    if resource not in RESOURCE_BRANCHES:
        raise HTTPException(404, "Unknown resource")
    module_name, branch = RESOURCE_BRANCHES[resource]
    with Session(engine) as session:
        _, data = load_module_data(session, module_name)
        for i, item in enumerate(data.get(branch, []) or []):
            if isinstance(item, dict) and record_id(item, i) == item_id:
                return item
    raise HTTPException(404, "Record not found")

@app.put("/api/v2/{resource}/{item_id}")
def v2_update(resource: str, item_id: str, payload: RecordPayload):
    if resource not in RESOURCE_BRANCHES:
        raise HTTPException(404, "Unknown resource")
    module_name, branch = RESOURCE_BRANCHES[resource]
    actor = payload.actor or "prototype-user"
    with Session(engine) as session:
        row, data = load_module_data(session, module_name)
        items = list(data.get(branch, []) or [])
        found = False
        for i, item in enumerate(items):
            if isinstance(item, dict) and record_id(item, i) == item_id:
                merged = dict(item); merged.update(payload.data); items[i] = merged; found = True; break
        if not found:
            raise HTTPException(404, "Record not found")
        data[branch] = items
        save_module_data(session, row, data, actor, f"{resource} updated", item_id)
        session.commit()
        return {"item": items[i], "revision": row.revision}

@app.delete("/api/v2/{resource}/{item_id}")
def v2_delete(resource: str, item_id: str, actor: str = "prototype-user"):
    if resource not in RESOURCE_BRANCHES:
        raise HTTPException(404, "Unknown resource")
    module_name, branch = RESOURCE_BRANCHES[resource]
    with Session(engine) as session:
        row, data = load_module_data(session, module_name)
        items = list(data.get(branch, []) or [])
        kept = [item for i, item in enumerate(items) if not (isinstance(item, dict) and record_id(item, i) == item_id)]
        if len(kept) == len(items):
            raise HTTPException(404, "Record not found")
        data[branch] = kept
        save_module_data(session, row, data, actor, f"{resource} deleted", item_id)
        session.commit()
        return {"ok": True, "revision": row.revision}

@app.post("/api/v2/requests/{request_id}/approval-decisions")
def record_approval_decision(request_id: str, payload: ApprovalDecisionPayload):
    with Session(engine) as session:
        row, data = load_module_data(session, "requests")
        requests = list(data.get("requests", []) or [])
        req = next((r for r in requests if isinstance(r, dict) and str(r.get("id")) == request_id), None)
        if not req:
            raise HTTPException(404, "Request not found")
        if req.get("status") != "Pending Approval":
            raise HTTPException(409, "Only a request awaiting approval can receive a decision")
        approvals = list(req.get("approvals", []) or [])

        # Business decisions are strictly sequential. Existing demo data may
        # still contain several rows marked Pending, so normalize it before
        # validating the actor and again after an approval advances the route.
        def prepare_sequence() -> dict[str, Any] | None:
            current = None
            for approval in approvals:
                state = approval.get("state")
                if state == "Approved" or state in {"Rejected", "Returned", "Invalidated"}:
                    continue
                if current is None:
                    approval["state"] = "Pending"
                    current = approval
                elif state in {"Pending", "Not Started", "Not started", "Queued"}:
                    approval["state"] = "Not started"
            return current

        current = prepare_sequence()
        if current is None:
            raise HTTPException(409, "This approval route has no pending decision")
        if current.get("role") != payload.role:
            raise HTTPException(
                409,
                f"Approval is currently awaiting {current.get('role')}; later approvers are locked",
            )
        existing = current
        decision = {"role": payload.role, "who": payload.actor, "state": payload.decision, "ts": now_iso(), "note": payload.note or ""}
        existing.update(decision)
        req["approvals"] = approvals
        if payload.decision == "Approved":
            next_approval = prepare_sequence()
            if next_approval is None and approvals and all(a.get("state") == "Approved" for a in approvals):
                req["status"] = "Approved"
        elif payload.decision == "Rejected":
            req["status"] = "Rejected"
            for approval in approvals:
                if approval is not existing and approval.get("state") in {"Pending", "Not Started", "Not started"}:
                    approval["state"] = "Invalidated"
        elif payload.decision == "Returned":
            req["status"] = "Returned for Correction"
            for approval in approvals:
                if approval is not existing and approval.get("state") in {"Pending", "Not Started", "Not started"}:
                    approval["state"] = "Invalidated"
        data["requests"] = requests
        save_module_data(session, row, data, payload.actor, "Approval decision recorded", request_id)
        session.commit()
        return {"request": req, "revision": row.revision}

@app.post("/api/v2/releases/{request_id}/tests/{test_id}/execute")
def execute_test_case(request_id: str, test_id: str, payload: TestExecutionPayload):
    with Session(engine) as session:
        row, data = load_module_data(session, "testing_release")
        pipeline = dict(data.get("releasePipeline", {}) or {})
        rel = pipeline.get(request_id)
        if not isinstance(rel, dict):
            raise HTTPException(404, "Release not found")
        cases = rel.get("testCases") or rel.get("tests") or []
        case = next((t for t in cases if str(t.get("id")) == test_id), None)
        if not case:
            raise HTTPException(404, "Test case not found")
        executions = list(case.get("executions", []) or [])
        execution = {"id": f"EXEC-{len(executions)+1:03d}", "tester": payload.actor, "actual": payload.actual_result,
                     "result": payload.result, "evidence": payload.evidence, "reference": payload.reference,
                     "comments": payload.comments, "executedAt": now_iso()}
        executions.append(execution)
        case["executions"] = executions; case["actual"] = payload.actual_result; case["res"] = payload.result
        rel["testCases"] = cases; pipeline[request_id] = rel; data["releasePipeline"] = pipeline
        save_module_data(session, row, data, payload.actor, "Test execution recorded", f"{request_id}:{test_id}")
        session.commit()
        return {"execution": execution, "test_case": case, "revision": row.revision}


def records(entity_type: str):
    with Session(engine) as session:
        rows = session.scalars(select(EntityRecord).where(EntityRecord.entity_type == entity_type)).all()
        return [json.loads(r.raw_json) for r in rows]

@app.get("/api/products")
def products(): return records("products")
@app.get("/api/templates")
def templates(): return records("templates")
@app.get("/api/requests")
def requests(): return records("requests")
@app.get("/api/users")
def users(): return records("users")
@app.get("/api/exceptions")
def exceptions(): return records("exceptions")
@app.get("/api/releases")
def releases(): return records("release_packages")

@app.get("/api/{entity_type}/{entity_id}")
def entity_detail(entity_type: str, entity_id: str):
    if entity_type not in {"products", "templates", "requests", "users", "exceptions", "release_packages", "dependencies", "systems", "retirements"}:
        raise HTTPException(404, "Unknown entity type")
    with Session(engine) as session:
        row = session.get(EntityRecord, {"entity_type": entity_type, "entity_id": entity_id})
        if not row:
            raise HTTPException(404, "Record not found")
        return json.loads(row.raw_json)

@app.get("/api/audit")
def audit_events(limit: int = 200):
    with Session(engine) as session:
        rows = session.scalars(select(AuditEvent).order_by(AuditEvent.id.desc()).limit(min(limit, 1000))).all()
        return [json.loads(r.raw_json) for r in rows]

@app.post("/api/auth/login")
def login(payload: LoginPayload):
    email = payload.email.strip().lower()
    for user in records("users"):
        if str(user.get("email", "")).lower() == email:
            if (user.get("accessStatus") or user.get("state")) not in ("Active", "active"):
                raise HTTPException(403, "User is not active")
            return {"user": user}
    raise HTTPException(404, "Provisioned PLM user not found")

EXPECTED_TEST_HEADERS = [
    "Scenario ID", "Scenario", "Objective", "System / Component", "Test Type", "Priority",
    "Mandatory", "Preconditions", "Test Data", "Test Steps", "Expected Result"
]

def parse_upload(file_name: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        return [dict(r) for r in csv.DictReader(io.StringIO(text))]
    if suffix in (".xlsx", ".xls"):
        if suffix == ".xls":
            # openpyxl cannot parse legacy binary XLS. The frontend accepts XLS
            # for compatibility, but the demo API asks users to save as XLSX.
            raise HTTPException(400, "Legacy .xls is not supported by this demo backend. Save the workbook as .xlsx and retry.")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(x).strip() if x is not None else "" for x in values[0]]
        return [dict(zip(headers, row)) for row in values[1:] if any(v is not None and str(v).strip() for v in row)]
    raise HTTPException(400, "Upload a CSV or XLSX file")

def store_import(import_type: str, file_name: str, rows: list[dict[str, Any]], valid: int, errors: list[dict[str, Any]]) -> str:
    batch_id = f"IMP-{datetime.now().strftime('%Y%m%d%H%M%S')}-{hashlib.sha1(file_name.encode()).hexdigest()[:4].upper()}"
    with Session(engine) as session:
        session.add(ImportBatch(
            batch_id=batch_id, import_type=import_type, file_name=file_name,
            row_count=len(rows), valid_count=valid, error_count=len(errors),
            status="Validated" if not errors else "Validation Errors",
            raw_json=json.dumps({"rows": rows, "errors": errors}, ensure_ascii=False, default=str),
        ))
        session.commit()
    return batch_id

@app.post("/api/import/test-cases")
async def import_test_cases(file: UploadFile = File(...)):
    content = await file.read()
    rows = parse_upload(file.filename or "upload", content)
    errors = []
    valid_rows = []
    for idx, row in enumerate(rows, start=2):
        missing = [h for h in ("Scenario", "Objective", "System / Component", "Expected Result") if not str(row.get(h, "") or "").strip()]
        if missing:
            errors.append({"row": idx, "errors": [f"Missing {m}" for m in missing]})
        else:
            valid_rows.append(row)
    batch_id = store_import("test_cases", file.filename or "upload", rows, len(valid_rows), errors)
    return {"batch_id": batch_id, "rows": rows, "valid_rows": valid_rows, "errors": errors, "expected_headers": EXPECTED_TEST_HEADERS}

@app.post("/api/import/estate")
async def import_estate(file: UploadFile = File(...)):
    content = await file.read()
    rows = parse_upload(file.filename or "upload", content)
    errors = []
    warnings = []
    required_any = ["Plan/Product Name", "Product Name", "name"]
    for idx, row in enumerate(rows, start=2):
        if not any(str(row.get(h, "") or "").strip() for h in required_any):
            errors.append({"row": idx, "errors": ["Plan/Product Name is required"]})
        if not (row.get("Product Domain") or row.get("Line of Business") or row.get("Domain")):
            warnings.append({"row": idx, "warning": "Product Domain mapping required"})
    batch_id = store_import("existing_estate", file.filename or "upload", rows, len(rows)-len(errors), errors)
    return {"batch_id": batch_id, "rows": rows, "errors": errors, "warnings": warnings}

@app.get("/api/imports")
def list_imports():
    with Session(engine) as session:
        rows = session.scalars(select(ImportBatch).order_by(ImportBatch.id.desc())).all()
        return [{
            "batch_id": r.batch_id, "type": r.import_type, "file_name": r.file_name,
            "rows": r.row_count, "valid": r.valid_count, "errors": r.error_count,
            "status": r.status, "created_at": r.created_at.isoformat() if r.created_at else None,
        } for r in rows]

@app.post("/api/mock/{system}/deploy")
def mock_deploy(system: str, payload: MockDeployPayload):
    result = payload.result.lower()
    ok = result in ("success", "ok", "passed")
    ref = f"{system.upper()}-{datetime.now().strftime('%Y%m%d%H%M%S')}"
    return {
        "system": system,
        "request_id": payload.request_id,
        "package_id": payload.package_id,
        "status": "Success" if ok else "Failed",
        "external_reference": ref,
        "message": payload.message or ("Configuration acknowledged" if ok else "Simulated downstream failure"),
        "timestamp": now_iso(),
    }

@app.post("/api/attachments/{entity_type}/{entity_id}")
async def upload_attachment(entity_type: str, entity_id: str, file: UploadFile = File(...)):
    safe_name = Path(file.filename or "attachment.bin").name
    folder = UPLOAD_DIR / entity_type / entity_id
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    path = folder / f"{stamp}_{safe_name}"
    with path.open("wb") as out:
        shutil.copyfileobj(file.file, out)
    with Session(engine) as session:
        rec = Attachment(entity_type=entity_type, entity_id=entity_id, file_name=safe_name, stored_path=str(path.relative_to(ROOT)), mime_type=file.content_type)
        session.add(rec)
        session.commit()
        session.refresh(rec)
    return {"id": rec.id, "file_name": safe_name, "path": rec.stored_path}

@app.get("/api/attachments/{attachment_id}")
def get_attachment(attachment_id: int):
    with Session(engine) as session:
        rec = session.get(Attachment, attachment_id)
        if not rec:
            raise HTTPException(404, "Attachment not found")
        path = ROOT / rec.stored_path
        return FileResponse(path, media_type=rec.mime_type, filename=rec.file_name)

# Static frontend is mounted last so /api routes win.
app.mount("/", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="frontend")
